from __future__ import annotations

from openpyxl import load_workbook

from proposal_xtraction import excel
from proposal_xtraction.validation import checar

AMARELO = "00FFF2A8"


def celulas(ws):
    return {c.value: c for row in ws.iter_rows() for c in row if c.value is not None}


def test_planilha_tem_cabecalho_e_itens(tmp_path, proposta_ok):
    destino = excel.build(proposta_ok, [], tmp_path / "p.xlsx")
    ws = load_workbook(destino).active

    assert ws.title == "Itens"
    valores = {c.value for row in ws.iter_rows() for c in row}
    assert "Metalúrgica Andrade Ltda" in valores
    assert "Chapa de aço 2mm" in valores
    assert "Perfil U 50x25" in valores
    assert 2000.0 in valores  # total geral


def test_itens_ficam_na_ordem_do_documento(tmp_path, proposta_ok):
    ws = load_workbook(excel.build(proposta_ok, [], tmp_path / "p.xlsx")).active
    descricoes = [
        row[2].value for row in ws.iter_rows() if isinstance(row[0].value, int) and row[2].value
    ]
    assert descricoes == ["Chapa de aço 2mm", "Perfil U 50x25"]


def test_aviso_de_linha_destaca_a_celula(tmp_path, proposta_ok):
    proposta_ok.itens[0].valor_total = 1400.0
    avisos = checar(proposta_ok)
    ws = load_workbook(excel.build(proposta_ok, avisos, tmp_path / "p.xlsx")).active

    destacadas = [
        c for row in ws.iter_rows() for c in row if c.fill.fgColor.rgb == AMARELO and c.comment
    ]
    assert destacadas, "esperava ao menos uma célula amarela com comentário"
    assert any("1.500,00" in c.comment.text for c in destacadas)


def test_aviso_de_cabecalho_destaca_o_campo(tmp_path, proposta_ok):
    proposta_ok.fornecedor = None
    avisos = checar(proposta_ok)
    ws = load_workbook(excel.build(proposta_ok, avisos, tmp_path / "p.xlsx")).active

    marcadas = [
        c for row in ws.iter_rows() for c in row if c.fill.fgColor.rgb == AMARELO and c.comment
    ]
    assert any("Fornecedor" in c.comment.text for c in marcadas)


def test_valores_monetarios_tem_formato_de_moeda(tmp_path, proposta_ok):
    ws = load_workbook(excel.build(proposta_ok, [], tmp_path / "p.xlsx")).active
    total = next(c for row in ws.iter_rows() for c in row if c.value == 1500.0)
    assert "R$" in total.number_format


def test_proposta_sem_itens_ainda_gera_planilha(tmp_path):
    from proposal_xtraction.models import Proposta

    proposta = Proposta()
    avisos = checar(proposta)
    ws = load_workbook(excel.build(proposta, avisos, tmp_path / "vazia.xlsx")).active
    textos = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("nenhum item" in t for t in textos)
