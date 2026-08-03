"""Geração da planilha a partir da proposta extraída."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Proposta
from .validation import Aviso

MOEDA_FMT = '"R$" #,##0.00'
DESTAQUE = PatternFill("solid", fgColor="FFF2A8")
CABECALHO_TABELA = PatternFill("solid", fgColor="D9D9D9")
NEGRITO = Font(bold=True)
BORDA_FINA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

COLUNAS = [
    ("Item", 8),
    ("Código", 16),
    ("Descrição", 55),
    ("Qtd.", 10),
    ("Un.", 8),
    ("Preço unit.", 14),
    ("Desconto", 14),
    ("Total", 16),
]

CAMPOS_CABECALHO = [
    ("Fornecedor", "fornecedor"),
    ("CNPJ/CPF", "cnpj"),
    ("Nº da proposta", "numero_proposta"),
    ("Data de emissão", "data_emissao"),
    ("Validade", "validade"),
    ("Condições de pagamento", "condicoes_pagamento"),
    ("Prazo de entrega", "prazo_entrega"),
    ("Moeda", "moeda"),
]

CAMPOS_TOTAIS = [
    ("Subtotal", "subtotal"),
    ("Frete", "frete"),
    ("Impostos", "impostos"),
    ("Total geral", "total_geral"),
]

# Campo da proposta -> coluna correspondente na tabela de itens (1-based).
COLUNA_POR_CAMPO = {
    "codigo": 2,
    "descricao": 3,
    "quantidade": 4,
    "unidade": 5,
    "preco_unitario": 6,
    "desconto": 7,
    "valor_total": 8,
}


def _marcar(cell, mensagens: list[str]) -> None:
    cell.fill = DESTAQUE
    texto = "\n".join(mensagens)
    cell.comment = Comment(texto, "Proposal-Xtraction", width=320, height=110)


def build(proposta: Proposta, avisos: list[Aviso], destino: Path) -> Path:
    """Escreve o .xlsx em `destino` e devolve o caminho."""
    por_linha: dict[int, list[Aviso]] = defaultdict(list)
    por_campo: dict[str, list[Aviso]] = defaultdict(list)
    for aviso in avisos:
        if aviso.linha:
            por_linha[aviso.linha].append(aviso)
        else:
            por_campo[aviso.campo].append(aviso)

    wb = Workbook()
    ws = wb.active
    ws.title = "Itens"

    linha = 1
    ws.cell(row=linha, column=1, value="Proposta comercial").font = Font(bold=True, size=14)
    linha += 2

    for rotulo, campo in CAMPOS_CABECALHO:
        ws.cell(row=linha, column=1, value=rotulo).font = NEGRITO
        valor = getattr(proposta, campo)
        cell = ws.cell(row=linha, column=2, value=valor if valor is not None else "—")
        cell.alignment = Alignment(horizontal="left")
        if campo in por_campo:
            _marcar(cell, [a.mensagem for a in por_campo[campo]])
        linha += 1

    linha += 1
    inicio_tabela = linha
    for indice, (titulo, largura) in enumerate(COLUNAS, start=1):
        cell = ws.cell(row=linha, column=indice, value=titulo)
        cell.font = NEGRITO
        cell.fill = CABECALHO_TABELA
        cell.border = BORDA_FINA
        ws.column_dimensions[get_column_letter(indice)].width = largura
    linha += 1

    for numero, item in enumerate(proposta.itens, start=1):
        valores = [
            numero,
            item.codigo,
            item.descricao,
            item.quantidade,
            item.unidade,
            item.preco_unitario,
            item.desconto,
            item.valor_total,
        ]
        for coluna, valor in enumerate(valores, start=1):
            cell = ws.cell(row=linha, column=coluna, value=valor)
            cell.border = BORDA_FINA
            if coluna in (6, 7, 8):
                cell.number_format = MOEDA_FMT
        ws.cell(row=linha, column=3).alignment = Alignment(wrap_text=True, vertical="top")

        for aviso in por_linha.get(numero, []):
            coluna = COLUNA_POR_CAMPO.get(aviso.campo, 8)
            _marcar(ws.cell(row=linha, column=coluna), [aviso.mensagem])
        linha += 1

    fim_tabela = linha - 1
    if proposta.itens:
        ws.auto_filter.ref = f"A{inicio_tabela}:H{fim_tabela}"
        ws.freeze_panes = ws.cell(row=inicio_tabela + 1, column=1)

    linha += 1
    for rotulo, campo in CAMPOS_TOTAIS:
        valor = getattr(proposta, campo)
        if valor is None and campo not in ("subtotal", "total_geral"):
            continue
        ws.cell(row=linha, column=7, value=rotulo).font = NEGRITO
        cell = ws.cell(row=linha, column=8, value=valor)
        cell.number_format = MOEDA_FMT
        cell.font = NEGRITO if campo == "total_geral" else Font()
        if campo in por_campo:
            _marcar(cell, [a.mensagem for a in por_campo[campo]])
        linha += 1

    if proposta.observacoes:
        linha += 1
        ws.cell(row=linha, column=1, value="Observações").font = NEGRITO
        ws.cell(row=linha + 1, column=1, value=proposta.observacoes).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        linha += 2

    gerais = [a for a in avisos if a.campo == "itens" and a.linha is None]
    if gerais:
        linha += 1
        ws.cell(row=linha, column=1, value="Avisos").font = NEGRITO
        for aviso in gerais:
            linha += 1
            cell = ws.cell(row=linha, column=1, value=aviso.mensagem)
            cell.fill = DESTAQUE

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino
